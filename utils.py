import re
import json
import pandas as pd
from docx import Document
import gspread
from google.oauth2.service_account import Credentials
import streamlit as st

SPREADSHEET_NAME = "QuanLy_DonDangKy_KHCN"
SHEET_NAME = "QUAN_LY_HO_SO"

LIST_LINH_VUC = [
    "Khoa học tự nhiên", "Khoa học kỹ thuật và công nghệ", "Khoa học y, dược",
    "Khoa học nông nghiệp", "Khoa học xã hội", "Khoa học nhân văn", "Công nghệ chiến lược"
]
LIST_LOAI_HINH = [
    "Nghiên cứu cơ bản", "Nghiên cứu ứng dụng", "Phát triển công nghệ", "Phát triển giải pháp xã hội"
]
LIST_HT_TRIEN_KHAI_1 = ["Nhiệm vụ", "Cụm nhiệm vụ", "Chuỗi nhiệm vụ"]
LIST_HT_TRIEN_KHAI_2 = [
    "Thực hiện theo hình thức liên kết", 
    "Thực hiện theo hình thức hợp tác công tư", 
    "Không thuộc 02 trường hợp trên"
]
LIST_HT_XET = ["Đặt hàng", "Tài trợ"]
LIST_PHAN_LOAI_CN = [
    "Nhân tài về KH, CN & ĐMST", 
    "Nhà khoa học trẻ / Kỹ sư trẻ tài năng", 
    "Không thuộc hai trường hợp trên"
]
LIST_HOC_HAM = ["Giáo sư", "Phó Giáo sư", "Tiến sĩ", "Thạc sĩ", "Cử nhân"]
LIST_TRANG_THAI = ["Đang xét", "Đang thực hiện", "Đã hoàn thành"]
LIST_NEN_TANG_SO = ["Đã cập nhật", "Chưa cập nhật"]

@st.cache_resource
def get_gspread_client():
    scopes = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]
    if "gcp_json" in st.secrets:
        json_str = st.secrets["gcp_json"]
        creds_dict = json.loads(json_str) if isinstance(json_str, str) else dict(json_str)
        creds = Credentials.from_service_account_info(creds_dict, scopes=scopes)
    elif "gcp_service_account" in st.secrets:
        creds_dict = dict(st.secrets["gcp_service_account"])
        creds = Credentials.from_service_account_info(creds_dict, scopes=scopes)
    else:
        creds = Credentials.from_service_account_file("credentials.json", scopes=scopes)
    return gspread.authorize(creds)

def get_worksheet():
    client = get_gspread_client()
    spreadsheet = client.open(SPREADSHEET_NAME)
    try:
        return spreadsheet.worksheet(SHEET_NAME)
    except Exception:
        return spreadsheet.sheet1

def is_checked(text):
    return any(sym in text for sym in ["☒", "☑", "■", "[x]", "[X]", "(x)", "(X)", "v", "V"])

def parse_bm09_word(doc_file):
    doc = Document(doc_file)
    full_paragraphs = [p.text for p in doc.paragraphs]
    table_texts = [" | ".join([cell.text.strip() for cell in row.cells]) for table in doc.tables for row in table.rows]
    all_lines = full_paragraphs + table_texts
    full_text = "\n".join(all_lines)

    extracted = {}
    match_ten = re.search(r"1\.\ Tên nhiệm vụ/cụm nhiệm vụ/chuỗi nhiệm vụ:\s*(.*?)(?=\n|Thuộc lĩnh vực:|$)", full_text, re.DOTALL)
    extracted["ten_nhiem_vu"] = match_ten.group(1).strip().strip("…").strip() if match_ten else ""

    extracted["linh_vuc"] = LIST_LINH_VUC[0]
    for lv in LIST_LINH_VUC:
        if any(lv.lower() in line.lower() and is_checked(line) for line in all_lines):
            extracted["linh_vuc"] = lv
            break

    extracted["loai_hinh"] = LIST_LOAI_HINH[0]
    for lh in LIST_LOAI_HINH:
        if any(lh.lower() in line.lower() and is_checked(line) for line in all_lines):
            extracted["loai_hinh"] = lh
            break

    extracted["ht_tk1"] = "Nhiệm vụ"
    extracted["ht_tk2"] = "Không thuộc 02 trường hợp trên"
    if "3.2. Cụm nhiệm vụ:" in full_text and is_checked(full_text.split("3.2. Cụm nhiệm vụ:")[1][:50]):
        extracted["ht_tk1"] = "Cụm nhiệm vụ"
    elif "3.3. Chuỗi nhiệm vụ:" in full_text and is_checked(full_text.split("3.3. Chuỗi nhiệm vụ:")[1][:50]):
        extracted["ht_tk1"] = "Chuỗi nhiệm vụ"

    if "Thực hiện theo hình thức liên kết" in full_text and is_checked(full_text.split("Thực hiện theo hình thức liên kết")[1][:30]):
        extracted["ht_tk2"] = "Thực hiện theo hình thức liên kết"
    elif "Thực hiện theo hình thức hợp tác công tư" in full_text and is_checked(full_text.split("Thực hiện theo hình thức hợp tác công tư")[1][:30]):
        extracted["ht_tk2"] = "Thực hiện theo hình thức hợp tác công tư"

    extracted["ht_xet"] = "Tài trợ" if ("Tài trợ" in full_text and is_checked(full_text.split("Tài trợ")[0][-10:])) else "Đặt hàng"

    match_tc = re.search(r"5\.\ Tổ chức chủ trì:\s*\n?Tên:\s*(.*?)\n", full_text)
    extracted["ten_tc"] = match_tc.group(1).strip().strip("…").strip() if match_tc else ""
    match_dc = re.search(r"Địa chỉ:\s*(.*?)\n", full_text)
    extracted["dc_tc"] = match_dc.group(1).strip().strip("…").strip() if match_dc else ""
    match_dt_tc = re.search(r"Điện thoại:\s*(.*?)(?=Email:|Website:|\n|$)", full_text)
    extracted["dt_tc"] = match_dt_tc.group(1).strip().strip("…").strip() if match_dt_tc else ""

    match_dd = re.search(r"6\.\ Người đại diện theo pháp luật\s*\n?Họ tên:\s*(.*?)(?=Chức vụ:|\n|$)", full_text)
    extracted["dd_pl"] = match_dd.group(1).strip().strip("…").strip() if match_dd else ""
    match_cv = re.search(r"Chức vụ:\s*(.*?)(?=\n|$)", full_text)
    extracted["chuc_vu_dd"] = match_cv.group(1).strip().strip("…").strip() if match_cv else ""

    extracted["tl_phap_ly"] = "Chưa cập nhật"
    extracted["tt_tc_de_xuat"] = "Chưa cập nhật"
    if "6.1. Tài liệu chứng minh" in full_text:
        part61 = full_text.split("6.1. Tài liệu chứng minh")[1].split("6.2.")[0]
        if "Đã cập nhật" in part61 and is_checked(part61.split("Đã cập nhật")[1][:20]):
            extracted["tl_phap_ly"] = "Đã cập nhật"
    if "6.2. Thông tin về tổ chức đề xuất" in full_text:
        part62 = full_text.split("6.2. Thông tin về tổ chức đề xuất")[1].split("7.")[0]
        if "Đã cập nhật" in part62 and is_checked(part62.split("Đã cập nhật")[1][:20]):
            extracted["tt_tc_de_xuat"] = "Đã cập nhật"

    match_cn = re.search(r"7\.\ Cá nhân đăng ký chủ nhiệm nhiệm vụ:.*?\nHọ tên:\s*(.*?)(?=Học hàm|\n|$)", full_text, re.DOTALL)
    extracted["ho_ten_cn"] = match_cn.group(1).strip().strip("…").strip() if match_cn else ""
    match_hh = re.search(r"Học hàm, học vị:\s*(.*?)(?=Đơn vị công tác|\n|$)", full_text)
    hh_str = match_hh.group(1).strip().strip("…").strip() if match_hh else ""
    extracted["hoc_ham_hoc_vi"] = next((hh for hh in LIST_HOC_HAM if hh.lower() in hh_str.lower()), "Tiến sĩ")

    match_dv = re.search(r"Đơn vị công tác:\s*(.*?)(?=Email:|\n|$)", full_text)
    extracted["don_vi_ct"] = match_dv.group(1).strip().strip("…").strip() if match_dv else ""
    match_email = re.search(r"Email:\s*(.*?)(?=Điện thoại:|\n|$)", full_text)
    extracted["email_cn"] = match_email.group(1).strip().strip("…").strip() if match_email else ""
    match_dt_cn = re.search(r"Điện thoại:\s*(.*?)(?=\n|Lý lịch|$)", full_text)
    extracted["dt_cn"] = match_dt_cn.group(1).strip().strip("…").strip() if match_dt_cn else ""

    extracted["phan_loai_cn"] = "Không thuộc hai trường hợp trên"
    if "Nhân tài về khoa học" in full_text and is_checked(full_text.split("Nhân tài về khoa học")[1][:30]):
        extracted["phan_loai_cn"] = "Nhân tài về KH, CN & ĐMST"
    elif "Nhà khoa học trẻ tài năng" in full_text and is_checked(full_text.split("Nhà khoa học trẻ tài năng")[1][:30]):
        extracted["phan_loai_cn"] = "Nhà khoa học trẻ / Kỹ sư trẻ tài năng"

    extracted["ly_lich_cn"] = "Chưa cập nhật"
    if "Lý lịch cá nhân đăng ký chủ nhiệm" in full_text:
        part_ll = full_text.split("Lý lịch cá nhân đăng ký chủ nhiệm")[1].split("8.")[0]
        if "Đã cập nhật" in part_ll and is_checked(part_ll.split("Đã cập nhật")[1][:20]):
            extracted["ly_lich_cn"] = "Đã cập nhật"

    match_kp = re.search(r"8\.\ Tổng kinh phí đề xuất\s*([\d\.,]+)", full_text)
    raw_kp = match_kp.group(1).replace(".", "").replace(",", "") if match_kp else "0"
    extracted["tong_kinh_phi"] = float(raw_kp) if raw_kp.isdigit() else 0.0

    match_nsnn_pct = re.search(r"Ngân sách nhà nước:\s*([\d\.,]+)\s*%", full_text)
    pct_nsnn = float(match_nsnn_pct.group(1).replace(",", ".")) if match_nsnn_pct else 0.0
    extracted["kinh_phi_nsnn"] = (extracted["tong_kinh_phi"] * pct_nsnn) / 100.0
    extracted["kinh_phi_ngoai_nsnn"] = max(0.0, extracted["tong_kinh_phi"] - extracted["kinh_phi_nsnn"])

    match_tg = re.search(r"9\.\ Thời gian thực hiện:\s*(\d+)", full_text)
    extracted["thoi_gian_th"] = int(match_tg.group(1)) if match_tg else 12

    return extracted

@st.cache_data(ttl=5)
def load_data():
    try:
        ws = get_worksheet()
        all_values = ws.get_all_values()
        if len(all_values) < 2:
            return pd.DataFrame()
        
        headers = all_values[1]
        data_rows = all_values[2:]
        df = pd.DataFrame(data_rows).iloc[:, :len(headers)]
        df.columns = headers
        df = df.dropna(how='all')
        df = df[df.astype(str).join(axis=1).str.strip() != ""]

        for col in ["STT", "Tổng kinh phí đề xuất (VNĐ)", "Kinh phí NSNN (VNĐ)", "Kinh phí Ngoài NSNN (VNĐ)", "Thời gian thực hiện (Tháng)"]:
            if col in df.columns:
                df[col] = pd.to_numeric(df[col].astype(str).str.replace(",", "").str.replace(".", ""), errors='coerce').fillna(0)

        for col in ["Tỷ lệ NSNN (%)", "Tỷ lệ Ngoài NSNN (%)"]:
            if col in df.columns:
                df[col] = pd.to_numeric(df[col].astype(str).str.replace("%", "").str.replace(",", "."), errors='coerce').fillna(0.0)

        return df
    except Exception as e:
        st.error(f"Lỗi đọc dữ liệu Google Sheets: {e}")
        return pd.DataFrame()

def save_new_record(new_row_dict):
    try:
        ws = get_worksheet()
        headers = ws.row_values(2)
        row_values = [str(new_row_dict.get(h, "")) for h in headers]
        ws.append_row(row_values)
        st.cache_data.clear()
        return True, "Thêm hồ sơ mới thành công lên Google Sheets!"
    except Exception as e:
        return False, f"Lỗi lưu dữ liệu: {str(e)}"